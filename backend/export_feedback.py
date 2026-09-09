"""
Export the collected user conversations into one spreadsheet, ONE TAB PER BOT.

Where the data comes from
-------------------------
Every chat is logged by the backend into the SQLite file ``logs.db`` (table
``chat_logs``) via the ``/log`` endpoint - see ``database.py``. That is the only
place user "feedback" is gathered today; there is no separate rating widget, so
"feedback" here means the full transcripts of what users asked and how each bot
answered.

What this script produces
-------------------------
``backend/exports/myri-feedback-<timestamp>.xlsx`` with:

  * a "Summary" tab  - one row per bot: sessions, messages, first/last activity
  * one tab per bot  - every message, grouped by conversation, in order

and, alongside it, ``backend/exports/csv/<bot>.csv`` (same data, one file per
bot) for quick import anywhere.

Turning it into a Google Sheet
------------------------------
Drag the .xlsx into Google Drive (or File -> Import in Sheets): Google converts
it to a native Sheet and keeps every tab. Re-run this script and re-upload
whenever you want fresh numbers.

Usage
-----
    ./venv/bin/python export_feedback.py
    ./venv/bin/python export_feedback.py --db logs.db --out exports
"""
from __future__ import annotations

import argparse
import csv
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BACKEND_DIR = Path(__file__).resolve().parent

# Every bot gets a tab even if it has no logs yet, so the sheet is always the
# full picture. id -> display title. Keep in sync with frontend/constants.ts.
BOTS: dict[str, str] = {
    "guardians-club": "The Lost Guardians' Club",
    "superhero-universe": "My Superhero Universe",
    "daily-dialogue": "Daily dialogue with NOOR",
    "journey-beliefs": "The Journey of Fundamental Beliefs",
    "my-compass": "My Compass",
    "better-me": "Better Me",
}

# Short, tidy Excel tab names (<=31 chars, no : \ / ? * [ ]). Full title lives
# in the Summary tab.
TAB_NAMES: dict[str, str] = {
    "guardians-club": "Guardians' Club",
    "superhero-universe": "Superhero Universe",
    "daily-dialogue": "Daily Dialogue (NOOR)",
    "journey-beliefs": "Journey & Beliefs",
    "my-compass": "My Compass",
    "better-me": "Better Me",
}

COLUMNS = ["Logged at", "Conversation", "Turn", "Who", "Message", "Sources"]

HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet_title(bot_id: str) -> str:
    # Excel tab names: <=31 chars, and none of : \ / ? * [ ]
    name = TAB_NAMES.get(bot_id) or BOTS.get(bot_id, bot_id)
    for ch in r':\/?*[]':
        name = name.replace(ch, " ")
    return name[:31].strip()


def _rows_for(conn: sqlite3.Connection, bot_id: str) -> list[dict]:
    cur = conn.execute(
        """
        SELECT session_id, timestamp, sender, message, sources
        FROM chat_logs
        WHERE chatbot_id = ?
        ORDER BY session_id, id
        """,
        (bot_id,),
    )
    return [
        {
            "session_id": r[0],
            "timestamp": r[1],
            "sender": r[2],
            "message": r[3] or "",
            "sources": r[4] or "",
        }
        for r in cur.fetchall()
    ]


def _pretty_sources(raw: str) -> str:
    if not raw:
        return ""
    try:
        import json

        val = json.loads(raw)
        if isinstance(val, list):
            return ", ".join(str(v) for v in val)
        return str(val)
    except (ValueError, TypeError):
        return raw


def _autosize(ws, max_widths: dict[int, int]) -> None:
    for idx, width in max_widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 80)


def _fill_sheet(ws, rows: list[dict]) -> None:
    ws.append(COLUMNS)
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"

    widths = {i: len(h) for i, h in enumerate(COLUMNS, start=1)}
    session_order: dict[str, int] = {}
    turn = 0
    last_session = None

    for r in rows:
        sid = r["session_id"]
        if sid not in session_order:
            session_order[sid] = len(session_order) + 1
        if sid != last_session:
            turn = 0
            last_session = sid
        turn += 1

        values = [
            r["timestamp"],
            f"#{session_order[sid]}",
            turn,
            "User" if r["sender"] == "user" else "Bot",
            r["message"],
            _pretty_sources(r["sources"]),
        ]
        ws.append(values)
        row_idx = ws.max_row
        ws.cell(row=row_idx, column=5).alignment = WRAP
        ws.cell(row=row_idx, column=6).alignment = WRAP
        for i, v in enumerate(values, start=1):
            widths[i] = max(widths[i], min(len(str(v)), 60))

    _autosize(ws, widths)


def _fill_summary(ws, conn: sqlite3.Connection) -> None:
    ws.append(["Bot", "Bot id", "Conversations", "Messages", "First activity", "Last activity"])
    for c in range(1, 7):
        ws.cell(row=1, column=c).fill = HEADER_FILL
        ws.cell(row=1, column=c).font = HEADER_FONT
    ws.freeze_panes = "A2"

    for bot_id, title in BOTS.items():
        row = conn.execute(
            """
            SELECT COUNT(DISTINCT session_id), COUNT(*), MIN(timestamp), MAX(timestamp)
            FROM chat_logs WHERE chatbot_id = ?
            """,
            (bot_id,),
        ).fetchone()
        ws.append([title, bot_id, row[0] or 0, row[1] or 0, row[2] or "-", row[3] or "-"])

    for i, w in enumerate([34, 20, 15, 12, 26, 26], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", default=str(BACKEND_DIR / "logs.db"), help="path to logs.db")
    ap.add_argument("--out", default=str(BACKEND_DIR / "exports"), help="output directory")
    args = ap.parse_args()

    db_path = Path(args.db)
    if not db_path.exists():
        raise SystemExit(f"No database at {db_path} - has anyone chatted yet?")

    out_dir = Path(args.out)
    csv_dir = out_dir / "csv"
    out_dir.mkdir(parents=True, exist_ok=True)
    csv_dir.mkdir(parents=True, exist_ok=True)

    conn = sqlite3.connect(str(db_path))

    # any bot ids present in the data but not in BOTS (renamed/new bots)
    known = set(BOTS)
    extra = [
        r[0]
        for r in conn.execute("SELECT DISTINCT chatbot_id FROM chat_logs").fetchall()
        if r[0] not in known
    ]
    for e in extra:
        BOTS[e] = e

    wb = Workbook()
    _fill_summary(wb.active, conn)
    wb.active.title = "Summary"

    total = 0
    for bot_id in BOTS:
        rows = _rows_for(conn, bot_id)
        total += len(rows)
        ws = wb.create_sheet(title=_sheet_title(bot_id))
        _fill_sheet(ws, rows)

        with open(csv_dir / f"{bot_id}.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(COLUMNS)
            sess: dict[str, int] = {}
            turn = 0
            last = None
            for r in rows:
                sid = r["session_id"]
                sess.setdefault(sid, len(sess) + 1)
                if sid != last:
                    turn, last = 0, sid
                turn += 1
                w.writerow([
                    r["timestamp"], f"#{sess[sid]}", turn,
                    "User" if r["sender"] == "user" else "Bot",
                    r["message"], _pretty_sources(r["sources"]),
                ])

    conn.close()

    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    xlsx_path = out_dir / f"myri-feedback-{stamp}.xlsx"
    wb.save(xlsx_path)

    print(f"Wrote {xlsx_path}")
    print(f"  {len(BOTS)} bot tabs, {total} messages total")
    print(f"  per-bot CSVs in {csv_dir}")


if __name__ == "__main__":
    main()
